#!/usr/bin/env python3
"""
Точки контактной сети/перегона, не связанные с сигналами: переезды, воздушные промежутки,
нейтральные вставки, обрывные места, блок-посты. Все пять файлов машинист прислал как
отдельные xlsx (один лист на файл, имя листа = тип точки), два разных формата:

  1. "Флаговый" формат (Переезды, Воздушный промежуток) — ОДНА физическая точка на строку:
     км, пк, затем "+"/пусто по каждому направлению — означает, нужно ли объявлять/учитывать
     эту точку при движении в эту сторону (не два разных места, одно и то же с отметкой,
     актуально ли оно для каждого направления). У переездов ещё есть "охраняемый" (+/пусто).

  2. "Парный" формат (Нейтральная вставка, Обрывное место, Блок-пост) — Четное/Нечетное это
     ДВЕ отдельные колонки км+пк (у нейтральной вставки — ещё и длина), т.е. в каждом
     направлении СВОЙ список точек, как и в остальных данных приложения (см.
     parse_brake_tests.py) — км+пк одной и той же физической точки, считанные по пикетажу
     каждого направления, физически могут не совпадать (разные пути на двухпутном участке,
     либо просто разный отсчёт пикетажа) — не пытаться сопоставлять строки между колонками
     по номеру строки, каждая колонка — самостоятельный список.

Направления в выходном JSON — "even"/"odd" (как во всех остальных ассетах), км+пк не
переводятся в метры здесь (это делает ChainageFormatter.toMeters() в Kotlin при загрузке).
"""
import sys
import json
import openpyxl

DIRECTION_MAP = {"чет": "even", "нечет": "odd"}


def parse_flagged_sheet(ws, guarded_col=None):
    """км, пк, [охраняемый], чет, нечет — одна точка на строку, +/пусто по направлениям."""
    rows = list(ws.iter_rows(values_only=True))
    header = next(r for r in rows if r[1] == "км")
    col_km = header.index("км")
    col_pk = header.index("пк")
    col_even = next(i for i, v in enumerate(header) if v and v.lower().startswith("чет"))
    col_odd = next(i for i, v in enumerate(header) if v and v.lower().startswith("нечет"))
    col_guard = header.index("охраняемый") if "охраняемый" in header else None

    entries = []
    for row in rows:
        if row[col_km] is None or row[col_km] == "км":
            continue
        km, pk = row[col_km], row[col_pk]
        if not isinstance(km, (int, float)):
            continue
        entries.append({
            "km": km,
            "pk": pk,
            "even": bool(row[col_even]),
            "odd": bool(row[col_odd]),
            **({"guarded": bool(row[col_guard])} if col_guard is not None else {}),
        })
    return entries


def parse_paired_sheet(ws, with_length=False):
    """Четное(км,пк[,длина]) | Нечетное(км,пк[,длина]) — раздельные списки по направлению."""
    rows = list(ws.iter_rows(values_only=True))
    header_row_idx = next(i for i, r in enumerate(rows) if "км" in r)
    label_row = rows[header_row_idx - 1]
    header = rows[header_row_idx]

    # Порядок колонок "Четное"/"Нечетное" не фиксирован (в разных файлах встречается и так,
    # и так) — определяем позиции обеих меток и берём диапазон колонок от метки до следующей
    # метки (или до конца строки для последней), а не полагаемся на то, что чётная всегда
    # первая.
    even_start = next(i for i, v in enumerate(label_row) if v and str(v).lower().startswith("чет"))
    odd_start = next(i for i, v in enumerate(label_row) if v and str(v).lower().startswith("нечет"))

    def cols_for(start, end):
        km_i = next(i for i in range(start, end) if header[i] == "км")
        pk_i = km_i + 1
        len_i = km_i + 2 if with_length else None
        return km_i, pk_i, len_i

    if even_start < odd_start:
        even_km, even_pk, even_len = cols_for(even_start, odd_start)
        odd_km, odd_pk, odd_len = cols_for(odd_start, len(header))
    else:
        odd_km, odd_pk, odd_len = cols_for(odd_start, even_start)
        even_km, even_pk, even_len = cols_for(even_start, len(header))

    entries = {"even": [], "odd": []}
    for row in rows[header_row_idx + 1:]:
        for direction, km_i, pk_i, len_i in (
            ("even", even_km, even_pk, even_len),
            ("odd", odd_km, odd_pk, odd_len),
        ):
            km = row[km_i]
            if not isinstance(km, (int, float)):
                continue
            entry = {"km": km, "pk": row[pk_i]}
            if with_length:
                entry["length_m"] = row[len_i]
            entries[direction].append(entry)
    return entries


def main():
    if len(sys.argv) < 4:
        print("usage: parse_catenary_points.py <format:flagged|paired> <xlsx> <out.json> [--length]")
        sys.exit(1)
    fmt, xlsx_path, out_path = sys.argv[1], sys.argv[2], sys.argv[3]
    with_length = "--length" in sys.argv[4:]

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb.worksheets[0]

    if fmt == "flagged":
        data = parse_flagged_sheet(ws)
        print(f"{ws.title}: {len(data)} точек")
    elif fmt == "paired":
        data = parse_paired_sheet(ws, with_length=with_length)
        print(f"{ws.title}: {len(data['even'])} чёт + {len(data['odd'])} неч точек")
    else:
        raise ValueError(f"unknown format {fmt!r}")

    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, separators=(",", ":"))
    print(f"Wrote {out_path}")


if __name__ == "__main__":
    main()
